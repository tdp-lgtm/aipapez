#!/usr/bin/env python3
"""Dump a .docx to readable text: paragraphs (with style), tables, AND real Word
footnotes, in document order.

NOTE (fixed 4 Jul 2026): python-docx's `.paragraphs`/body iteration does NOT include
footnote text (footnotes live in a separate word/footnotes.xml part). Earlier versions
of this script were therefore footnote-blind, so any reader (e.g. a referee agent) saw
the References list but not the footnotes that cite those works — reading them as
"in the bibliography but never cited." The footnote dump below fixes that; always read
the FOOTNOTES section as part of the manuscript."""
import re
import sys
import html
import zipfile
from docx import Document
from docx.document import Document as _Doc
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def dump_footnotes(path):
    """Print real Word footnotes (word/footnotes.xml), skipping the separator stubs."""
    try:
        z = zipfile.ZipFile(path)
        if "word/footnotes.xml" not in z.namelist():
            return 0
        xml = z.read("word/footnotes.xml").decode("utf-8", "replace")
    except Exception:
        return 0
    notes = []
    for m in re.finditer(r"<w:footnote\b([^>]*)>(.*?)</w:footnote>", xml, re.S):
        attrs, content = m.group(1), m.group(2)
        if "w:type=" in attrs:  # separator / continuationSeparator stubs
            continue
        text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", content))
        text = html.unescape(re.sub(r"\s+", " ", text)).strip()
        if text:
            notes.append(text)
    if notes:
        print("\n--- FOOTNOTES (part of the manuscript — read these) ---")
        for i, text in enumerate(notes, 1):
            print(f"[fn {i}] {text}")
        print("--- end footnotes ---")
    return len(notes)

def iter_block_items(parent):
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def main(path):
    doc = Document(path)
    n_para = 0
    n_tbl = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            n_para += 1
            style = block.style.name if block.style else ""
            text = block.text
            if text.strip() == "" and style.lower().startswith("normal"):
                continue
            tag = f"[{style}]" if style and style != "Normal" else ""
            print(f"{tag} {text}".rstrip())
        elif isinstance(block, Table):
            n_tbl += 1
            print(f"\n--- TABLE {n_tbl} ({len(block.rows)} rows x {len(block.columns)} cols) ---")
            for r in block.rows:
                cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
                print(" | ".join(cells))
            print("--- end table ---\n")
    n_fn = dump_footnotes(path)
    # word count of body text
    words = 0
    for p in doc.paragraphs:
        words += len(p.text.split())
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                words += len(c.text.split())
    sys.stderr.write(f"\n[stats] paragraphs={n_para} tables={n_tbl} "
                     f"footnotes={n_fn} approx_words={words}\n")

if __name__ == "__main__":
    main(sys.argv[1])
